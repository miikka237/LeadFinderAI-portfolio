"""Export results for the calling team.

Two outputs:
  * Leads  -- ACCEPT rows, ready for calling
  * Review -- REVIEW rows with the reason, for a human to check

REJECT rows stay in the database only. Excel (.xlsx) is written with
openpyxl; any other extension falls back to CSV files using only the
standard library.
"""

import csv
import os
from typing import Dict, List

COLUMNS = [
    "domain", "url", "company_name", "business_id", "business_type",
    "attribution_status", "source_query",
]
REVIEW_COLUMNS = COLUMNS + ["ai_status", "reason"]


def export_results(accept_rows: List[Dict], review_rows: List[Dict], path: str) -> List[str]:
    """Write the export and return the list of files created."""
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    if path.lower().endswith(".xlsx"):
        return [_export_xlsx(accept_rows, review_rows, path)]

    base, _ = os.path.splitext(path)
    leads_path, review_path = f"{base}_leads.csv", f"{base}_review.csv"
    _write_csv(leads_path, COLUMNS, accept_rows)
    _write_csv(review_path, REVIEW_COLUMNS, review_rows)
    return [leads_path, review_path]


def _write_csv(path: str, columns: List[str], rows: List[Dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _export_xlsx(accept_rows: List[Dict], review_rows: List[Dict], path: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [("Leads", COLUMNS, accept_rows), ("Review", REVIEW_COLUMNS, review_rows)]
    for i, (title, columns, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        ws.freeze_panes = "A2"
        for col_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 60)
    wb.save(path)
    return path
